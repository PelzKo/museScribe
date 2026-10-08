"""Generate the song folders of a show from one workbook + one show file.

  data/_<show>/show.yaml     constants: meta.yaml block, per-part defaults, instrument library
  data/_<show>/songs.xlsx    one sheet per song (sheet name = song id), one row per PDF
  data/_<show>/pdfs/         the full-book PDFs, copied into every song folder

Song sheet layout (sheets whose name starts with "_" are ignored):

  A1 "title"  B1 <song title>
  row 3 header: file | pages | skip | instruments | extra | notes
  file         PDF stem in pdfs/ (e.g. "Reed 1")
  pages        "11-16"; empty = part not in this song (no yaml, no PDF copied)
  skip         systems to leave out, "1-3" or "1,4" -> skip_systems: [1, 2, 3]
  instruments  "Clarinet, Flute": looked up in the instrument library. Several on a
               non-drums part -> split_instruments; a single one only adds its transpose.
               "Triangle (lower_staff)": keys in brackets for this song only.
  extra        raw YAML keys, "chord_names: false, lyrics: true"; "lyrics: null" leaves out a
               part default (the engine's own default applies)
  notes        ignored
"""
from __future__ import annotations

import difflib
import re
import shutil
from datetime import date, datetime
from pathlib import Path

import yaml

HEADER = ["file", "pages", "skip", "instruments", "extra", "notes"]
KEY_ORDER = ["pages", "type", "source", "transpose", "keep_staves", "skip_systems",
             "split_instruments", "doubling"]


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        raise ValueError(f"cell {v!r} was turned into a date by Excel; format the column as Text")
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def expand_range(spec: str) -> list[int]:
    """"1-3,7" -> [1, 2, 3, 7]."""
    out: list[int] = []
    for item in (x.strip() for x in spec.split(",")):
        if item:
            lo, _, hi = item.partition("-")
            out.extend(range(int(lo), int(hi or lo) + 1))
    return sorted(set(out))


def _fmt(s: str, ctx: dict):
    return s.format(**ctx) if isinstance(s, str) else s


def split_instruments_cell(cell: str) -> list[tuple[str, dict]]:
    """"Timpani, Triangle (lower_staff)" -> [("Timpani", {}), ("Triangle", {"lower_staff": True})].

    Brackets hold YAML keys for this song only; a bare key means true."""
    out = []
    for m in re.finditer(r"\s*([^,(]+?)\s*(?:\(([^)]*)\))?\s*(?:,|$)", cell):
        if not m.group(1):
            continue
        keys = yaml.safe_load("{" + m.group(2) + "}") if m.group(2) else {}
        out.append((m.group(1), {k: True if v is None else v for k, v in keys.items()}))
    return out


def _instrument(name: str, library: dict, own: dict | None = None):
    extra = {**(library.get(name) or {}), **(own or {})}
    return {"name": name, **extra} if extra else name


def part_meta(row: dict, show: dict, ctx: dict) -> dict:
    name = row["file"]
    meta = {"pages": row["pages"]}
    meta.update({k: _fmt(v, ctx) for k, v in (show.get("parts", {}).get(name) or {}).items()})
    if row["skip"]:
        meta["skip_systems"] = expand_range(row["skip"])
    entries = split_instruments_cell(row["instruments"])
    extra = yaml.safe_load("{" + row["extra"].replace("\n", ", ") + "}") if row["extra"] else {}
    library = show.get("instruments", {})
    drums = (extra.get("type") or meta.get("type")) == "drums"
    if any(own.get("lower_staff") for _, own in entries):
        # the sheet names this song's lower-staff instrument: it replaces the library's
        library = {k: {f: v for f, v in (x or {}).items() if f != "lower_staff"}
                   for k, x in library.items()}
    if len(entries) == 1 and not drums:
        name, own = entries[0]
        if t := {**(library.get(name) or {}), **own}.get("transpose"):
            meta["transpose"] = t
    elif entries:
        meta["instruments"] = [_instrument(n, library, own) for n, own in entries]
        if not drums and not extra.get("doubling"):
            meta["split_instruments"] = True
    meta.update(extra)
    meta = {k: v for k, v in meta.items() if v is not None}  # "lyrics: null" drops a part default
    order = KEY_ORDER + [k for k in meta if k not in KEY_ORDER and k != "instruments"] + ["instruments"]
    return {k: meta[k] for k in sorted(meta, key=order.index)}


def dump(meta: dict, header: str) -> str:
    """Top level in block style, values (instrument entries too) in flow style."""
    def flow(v) -> str:
        return yaml.safe_dump(v, default_flow_style=True, width=10**9, sort_keys=False).strip() \
            .removesuffix("\n...").removesuffix("...").strip()

    lines = [header]
    for k, v in meta.items():
        if isinstance(v, list) and any(isinstance(x, dict) for x in v):
            lines.append(f"{k}:")
            lines.extend(f"  - {flow(x)}" for x in v)
        elif k == "pages":
            lines.append(f'pages: "{v}"')
        else:
            lines.append(f"{k}: {flow(v)}")
    return "\n".join(lines) + "\n"


def read_workbook(path: Path) -> dict[str, dict]:
    """{song: {"title": str, "rows": [row dict]}} in sheet order."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    songs = {}
    for ws in wb.worksheets:
        if ws.title.startswith("_"):
            continue
        rows = list(ws.iter_rows(values_only=True))
        title = _cell(rows[0][1]) if rows and len(rows[0]) > 1 else ""
        hdr_i = next((i for i, r in enumerate(rows) if r and _cell(r[0]).lower() == "file"), None)
        if hdr_i is None:
            raise ValueError(f"sheet '{ws.title}': no header row starting with 'file'")
        cols = [_cell(c).lower() for c in rows[hdr_i]]
        out = []
        for r in rows[hdr_i + 1:]:
            try:
                row = {h: _cell(r[cols.index(h)]) if h in cols and cols.index(h) < len(r) else ""
                       for h in HEADER}
            except ValueError as e:
                raise ValueError(f"sheet '{ws.title}', row {_cell(r[0])!r}: {e}") from None
            if row["file"]:
                out.append(row)
        songs[ws.title] = {"title": title, "rows": out}
    return songs


def _write(path: Path, text: str, dry: bool) -> bool:
    old = path.read_text(encoding="utf-8") if path.exists() else None
    if old == text:
        return False
    if old is None:
        print(f"  new      {path.name}")
    else:
        print(f"  changed  {path.name}")
        for ln in difflib.unified_diff(old.splitlines(), text.splitlines(), lineterm="", n=0):
            if not ln.startswith(("---", "+++", "@@")):
                print(f"           {ln}")
    if not dry:
        path.write_text(text, encoding="utf-8")
    return True


def generate(show_file: Path, data_dir: Path, songs_only: list[str] | None = None,
             dry: bool = False) -> int:
    show_file = show_file.resolve()
    base = show_file.parent
    show = yaml.safe_load(show_file.read_text(encoding="utf-8")) or {}
    pdfs = base / show.get("pdfs", "pdfs")
    workbook = base / show.get("workbook", "songs.xlsx")
    songs = read_workbook(workbook)
    problems = 0
    for song, sheet in songs.items():
        if songs_only and song not in songs_only:
            continue
        ctx = {"song": song, "title": sheet["title"]}
        folder = data_dir / show.get("folder", "{song}").format(**ctx)
        print(f"{folder.name}  ({sheet['title']})")
        if not dry:
            folder.mkdir(parents=True, exist_ok=True)
        header = f"# generated by `omrbench gen-meta` from {workbook.name}, sheet '{song}' - edit there"
        meta = {k: _fmt(v, ctx) for k, v in (show.get("meta") or {}).items()}
        _write(folder / "meta.yaml", dump(meta, header), dry)
        for row in sheet["rows"]:
            name = row["file"]
            src, dst = pdfs / f"{name}.pdf", folder / f"{name}.pdf"
            if not row["pages"]:
                if dst.exists() or (folder / f"{name}.yaml").exists():
                    print(f"  WARN     {name}: no pages in the sheet, but {name}.pdf/.yaml is in "
                          f"the folder (would be transcribed whole); delete it")
                    problems += 1
                continue
            try:
                text = dump(part_meta(row, show, ctx), header)
            except Exception as e:
                print(f"  ERROR    {name}: {e}")
                problems += 1
                continue
            _write(folder / f"{name}.yaml", text, dry)
            if not src.exists():
                print(f"  ERROR    {name}: {src} missing")
                problems += 1
            elif not dst.exists() or dst.stat().st_size != src.stat().st_size:
                print(f"  copy     {name}.pdf")
                if not dry:
                    shutil.copy2(src, dst)
        listed = {r["file"] for r in sheet["rows"]}
        for pdf in sorted(folder.glob("*.pdf")) if folder.exists() else []:
            if pdf.stem not in listed:
                print(f"  WARN     {pdf.name} is not in the sheet (transcribed with meta.yaml only)")
                problems += 1
    if songs_only and (missing := set(songs_only) - set(songs)):
        print(f"WARN no sheet for: {', '.join(sorted(missing))}")
        problems += 1
    if dry:
        print("(dry run, nothing written)")
    return problems
