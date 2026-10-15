"""A show transcribed song by song from the same full-book parts: one workbook + one show file.

  data/<show>/show.yaml     constants: meta block, per-part defaults, instrument library
  data/<show>/songs.xlsx    one sheet per song (sheet name = song id), one row per PDF
  data/<show>/pdfs/         the full-book PDFs (read in place, never copied)
  data/<show>/truth/        ground-truth MuseScore files, linked from the workbook

Song sheet layout (sheets whose name starts with "_" are ignored):

  A1 "title"  B1 <song title>
  A2 "truth"  B2 <truth file, relative to the show folder>   (optional; empty = no benchmark)
  A3 "unfinished"  B3 <truth part names not hand-corrected yet, "Voice 1, Drum Kit"> (optional;
               bench leaves them out until they are removed here)
  header row: file | pages | skip | instruments | extra | truth | notes
  file         PDF stem in pdfs/ (e.g. "Reed 1")
  pages        "11-16"; empty = part not in this song
  skip         systems to leave out, "1-3" or "1,4" -> skip_systems: [1, 2, 3]
  instruments  "Clarinet, Flute": looked up in the instrument library. Several on a
               non-drums part -> split_instruments; a single one only adds its transpose.
               "Triangle (lower_staff)": keys in brackets for this song only.
  extra        raw YAML keys, "chord_names: false, lyrics: true"; "lyrics: null" leaves out a
               part default (the engine's own default applies)
  truth        part name(s) in the truth score this row is compared with (optional column)
  notes        ignored

Every key/value row above the header (column A = key) is a song setting; `title`, `truth` and
`unfinished` are used. Nothing is generated on disk: `load` builds every part's meta in memory,
`check` validates the workbook and shows what each part resolves to.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

import yaml

HEADER = ["file", "pages", "skip", "instruments", "extra", "truth", "notes"]
KEY_ORDER = ["pages", "type", "source", "transpose", "keep_staves", "skip_systems",
             "split_instruments", "doubling"]
SHOW_FILE = "show.yaml"


@dataclass
class Part:
    song: str
    name: str                  # PDF stem in pdfs/
    pdf: Path
    meta: dict                 # show meta + part meta (part keys win)
    truth: list[str] = field(default_factory=list)   # truth part name(s)
    instruments: list[str] = field(default_factory=list)


@dataclass
class Song:
    id: str
    title: str
    truth: Path | None
    parts: list[Part]
    errors: list[str] = field(default_factory=list)   # rows that could not be resolved
    unfinished: list[str] = field(default_factory=list)   # truth parts not hand-corrected yet


@dataclass
class Project:
    name: str
    dir: Path
    show: dict
    workbook: Path
    pdfs: Path
    songs: dict[str, Song]

    def select(self, songs: list[str] | None) -> tuple[list[Song], list[str]]:
        """Songs in sheet order (all if `songs` is empty), and the requested ids with no sheet."""
        missing = [s for s in songs or [] if s not in self.songs]
        return [s for k, s in self.songs.items() if not songs or k in songs], missing


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        raise ValueError(f"cell {v!r} was turned into a date by Excel; format the column as Text")
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def expand_range(spec: str) -> list[int]:
    """"1-3,7" -> [1, 2, 3, 7]."""
    out: list[int] = []
    for item in (x.strip() for x in spec.split(",")):
        if item:
            lo, _, hi = item.partition("-")
            out.extend(range(int(lo), int(hi or lo) + 1))
    return sorted(set(out))


def _fmt(s: str, ctx: dict):
    return s.format(**ctx) if isinstance(s, str) else s


def split_instruments_cell(cell: str) -> list[tuple[str, dict]]:
    """"Timpani, Triangle (lower_staff)" -> [("Timpani", {}), ("Triangle", {"lower_staff": True})].

    Brackets hold YAML keys for this song only; a bare key means true."""
    out = []
    for m in re.finditer(r"\s*([^,(]+?)\s*(?:\(([^)]*)\))?\s*(?:,|$)", cell):
        if not m.group(1):
            continue
        keys = yaml.safe_load("{" + m.group(2) + "}") if m.group(2) else {}
        out.append((m.group(1), {k: True if v is None else v for k, v in keys.items()}))
    return out


def _instrument(name: str, library: dict, own: dict | None = None):
    extra = {**(library.get(name) or {}), **(own or {})}
    return {"name": name, **extra} if extra else name


def part_meta(row: dict, show: dict, ctx: dict) -> dict:
    name = row["file"]
    meta = {"pages": row["pages"]}
    meta.update({k: _fmt(v, ctx) for k, v in (show.get("parts", {}).get(name) or {}).items()})
    if row["skip"]:
        meta["skip_systems"] = expand_range(row["skip"])
    entries = split_instruments_cell(row["instruments"])
    extra = yaml.safe_load("{" + row["extra"].replace("\n", ", ") + "}") if row["extra"] else {}
    library = show.get("instruments", {})
    drums = (extra.get("type") or meta.get("type")) == "drums"
    if any(own.get("lower_staff") for _, own in entries):
        # the sheet names this song's lower-staff instrument: it replaces the library's
        library = {k: {f: v for f, v in (x or {}).items() if f != "lower_staff"}
                   for k, x in library.items()}
    if len(entries) == 1 and not drums:
        name, own = entries[0]
        if t := {**(library.get(name) or {}), **own}.get("transpose"):
            meta["transpose"] = t
    elif entries:
        meta["instruments"] = [_instrument(n, library, own) for n, own in entries]
        if not drums and not extra.get("doubling"):
            meta["split_instruments"] = True
    meta.update(extra)
    meta = {k: v for k, v in meta.items() if v is not None}  # "lyrics: null" drops a part default
    order = KEY_ORDER + [k for k in meta if k not in KEY_ORDER and k != "instruments"] + ["instruments"]
    return {k: meta[k] for k in sorted(meta, key=order.index)}


def dump(meta: dict, header: str = "") -> str:
    """Top level in block style, values (instrument entries too) in flow style."""
    def flow(v) -> str:
        return yaml.safe_dump(v, default_flow_style=True, width=10**9, sort_keys=False).strip() \
            .removesuffix("\n...").removesuffix("...").strip()

    lines = [header] if header else []
    for k, v in meta.items():
        if isinstance(v, list) and any(isinstance(x, dict) for x in v):
            lines.append(f"{k}:")
            lines.extend(f"  - {flow(x)}" for x in v)
        elif k == "pages":
            lines.append(f'pages: "{v}"')
        else:
            lines.append(f"{k}: {flow(v)}")
    return "\n".join(lines) + "\n"


def read_workbook(path: Path) -> dict[str, dict]:
    """{song: {"settings": {key: value}, "rows": [row dict]}} in sheet order."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    songs = {}
    for ws in wb.worksheets:
        if ws.title.startswith("_"):
            continue
        rows = list(ws.iter_rows(values_only=True))
        hdr_i = next((i for i, r in enumerate(rows) if r and _cell(r[0]).lower() == "file"), None)
        if hdr_i is None:
            raise ValueError(f"sheet '{ws.title}': no header row starting with 'file'")
        settings = {_cell(r[0]).lower(): _cell(r[1]) if len(r) > 1 else ""
                    for r in rows[:hdr_i] if r and _cell(r[0])}
        cols = [_cell(c).lower() for c in rows[hdr_i]]
        out = []
        for r in rows[hdr_i + 1:]:
            try:
                row = {h: _cell(r[cols.index(h)]) if h in cols and cols.index(h) < len(r) else ""
                       for h in HEADER}
            except ValueError as e:
                raise ValueError(f"sheet '{ws.title}', row {_cell(r[0])!r}: {e}") from None
            if row["file"]:
                out.append(row)
        songs[ws.title] = {"settings": settings, "rows": out}
    return songs


def find(data_dir: Path, name: str) -> Path | None:
    """The show folder for a name ("anastasia") or a path (folder or its show.yaml)."""
    for p in (Path(name), Path(data_dir) / name):
        if p.is_file() and p.name == SHOW_FILE:
            return p.parent
        if (p / SHOW_FILE).is_file():
            return p
    return None


def projects(data_dir: Path) -> list[Path]:
    return sorted(p for p in Path(data_dir).iterdir() if (p / SHOW_FILE).is_file()) \
        if Path(data_dir).is_dir() else []


def load(folder: Path) -> Project:
    folder = Path(folder).resolve()
    show = yaml.safe_load((folder / SHOW_FILE).read_text(encoding="utf-8")) or {}
    pdfs = folder / show.get("pdfs", "pdfs")
    workbook = folder / show.get("workbook", "songs.xlsx")
    songs = {}
    for sid, sheet in read_workbook(workbook).items():
        st = sheet["settings"]
        ctx = {"song": sid, "title": st.get("title", "")}
        folder_meta = {k: _fmt(v, ctx) for k, v in (show.get("meta") or {}).items()}
        song = Song(sid, ctx["title"], folder / st["truth"] if st.get("truth") else None, [],
                    unfinished=[x.strip() for x in st.get("unfinished", "").split(",") if x.strip()])
        for row in sheet["rows"]:
            if not row["pages"]:
                continue
            try:
                meta = {**folder_meta, **part_meta(row, show, ctx)}
                if meta.get("type") == "drums":
                    meta["show"] = folder.name     # the pipeline looks up the show's drum legend
            except Exception as e:  # noqa: BLE001 - one bad row never hides the others
                song.errors.append(f"{row['file']}: {type(e).__name__}: {e}")
                continue
            song.parts.append(Part(sid, row["file"], pdfs / f"{row['file']}.pdf", meta,
                                   [x.strip() for x in row["truth"].split(",") if x.strip()],
                                   [n for n, _ in split_instruments_cell(row["instruments"])]))
        songs[sid] = song
    return Project(folder.name, folder, show, workbook, pdfs, songs)


def truth_names(part: Part) -> list[str]:
    """The truth part names a row's `truth` cell refers to ("Out=Truth" -> Truth; "-" = none)."""
    names = [e.partition("=")[2].strip() if "=" in e else e.strip() for e in part.truth]
    return [n for n in names if n and n != "-"]


def _suggest_truth(proj: Project, part: Part, tparts: dict[str, int]) -> str | None:
    """A `truth` cell for a row without one: the cell the same PDF has in another song if all its
    names are in this truth score, else the truth part names closest to the row's name /
    instruments (truth parts with notes only)."""
    import difflib

    names = {" ".join(n.lower().split()): n for n, c in tparts.items() if c}
    for other in proj.songs.values():
        for p in other.parts:
            if p.name == part.name and p.truth and all(
                    n == "-" or " ".join(n.lower().split()) in names for n in truth_names(p)):
                return ", ".join(p.truth)

    def close(word: str) -> str:
        hit = difflib.get_close_matches(word.lower(), list(names), n=1, cutoff=0.5)
        return names[hit[0]] if hit else "-"
    picks = [close(i) for i in part.instruments] if len(part.instruments) > 1 else [close(part.name)]
    return ", ".join(picks) if any(p != "-" for p in picks) else None


def check(proj: Project, songs_only: list[str] | None = None, verbose: bool = False,
          export: Path | None = None, truth_cache: Path | None = None,
          musescore: str | None = None) -> int:
    """Validate the workbook and print what every part resolves to; return the problem count.

    export: also write <dir>/<song>/<part>.yaml files (for inspection only).
    truth_cache: where truth scores are converted to (results/cache); the `truth` names are
    then checked against the score's part names."""
    from .rasterize import pdf_pages
    from .transcribe import parse_pages

    songs, missing = proj.select(songs_only)
    problems = 0

    def problem(msg: str):
        nonlocal problems
        problems += 1
        print(f"  {msg}")

    library = proj.show.get("instruments") or {}
    page_counts: dict[Path, int] = {}
    for song in songs:
        truth = f"   truth: {song.truth.name}" if song.truth else ""
        print(f"{proj.name} {song.id}  ({song.title}){truth}")
        tparts: dict[str, int] | None = None   # truth part name -> note count
        if song.truth and not song.truth.exists():
            problem(f"ERROR    truth file {song.truth} missing")
        elif song.truth and truth_cache is not None:
            from . import musicxml
            from . import truth as truth_mod
            from .metrics import notef1
            try:
                troot = musicxml.load(truth_mod.convert(song.truth, truth_cache, musescore))
                tparts = {n: len(notef1.notes(p)) for n, p in truth_mod.parts(troot)}
            except Exception as e:  # noqa: BLE001
                problem(f"ERROR    truth file {song.truth.name}: {e}")
        for e in song.errors:
            problem(f"ERROR    {e}")
        if song.unfinished:
            print(f"  note     unfinished truth parts (not scored): {', '.join(song.unfinished)}")
            if tparts is not None:
                known_t = {" ".join(n.lower().split()) for n in tparts}
                for n in song.unfinished:
                    if " ".join(n.lower().split()) not in known_t:
                        problem(f"ERROR    unfinished: no part '{n}' in {song.truth.name}")
        for part in song.parts:
            if not part.pdf.exists():
                problem(f"ERROR    {part.name}: {part.pdf} missing")
                continue
            if part.pdf not in page_counts:
                page_counts[part.pdf] = pdf_pages(part.pdf)
            try:
                n = len(parse_pages(part.meta["pages"], page_counts[part.pdf]))
            except ValueError as e:
                problem(f"ERROR    {part.name}: {e} (PDF has {page_counts[part.pdf]} pages)")
                continue
            unknown = [i for i in part.instruments if i not in library]
            if song.truth and not part.truth:
                hint = _suggest_truth(proj, part, tparts) if tparts else None
                problem(f"WARN     {part.name}: no `truth` part name (not benchmarked)"
                        + (f"; suggest: {hint}" if hint else ""))
            if tparts is not None:
                known = {" ".join(n.lower().split()) for n in tparts}
                for tn in truth_names(part):
                    if " ".join(tn.lower().split()) not in known:
                        problem(f"ERROR    {part.name}: no part '{tn}' in {song.truth.name}")
            info = f"{n:>2} pages ({part.meta['pages']})"
            if part.instruments:
                info += f"  [{', '.join(part.instruments)}]"
            if unknown:   # fine for instruments in C; a typo would lose its transposition
                info += f"  (not in show.yaml: {', '.join(unknown)})"
            if part.truth:
                info += f"  -> truth {', '.join(part.truth)}"
            print(f"  {part.name:<18} {info}")
            if verbose:
                for ln in dump(part.meta).splitlines():
                    print(f"      {ln}")
            if export:
                out = export / song.id / f"{part.name}.yaml"
                out.parent.mkdir(parents=True, exist_ok=True)
                out.write_text(dump(part.meta, f"# from {proj.workbook.name}, sheet '{song.id}'"),
                               encoding="utf-8")
        if tparts is not None:
            named = {" ".join(n.lower().split()) for p in song.parts for n in truth_names(p)}
            named |= {" ".join(n.lower().split()) for n in song.unfinished}
            if rest := [f"{n} ({c})" for n, c in tparts.items()
                        if c and " ".join(n.lower().split()) not in named]:
                print(f"  note     truth parts with notes no row names (notes): {', '.join(rest)}")
    if missing:
        problem(f"WARN no sheet for: {', '.join(missing)}")
    if export:
        print(f"part yamls written to {export}")
    return problems

